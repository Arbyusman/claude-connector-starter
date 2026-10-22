"""docimage MCP server: read and insert images in Word (.docx) and PDF files.

Tools
  list_images        - list the images inside a .docx/.pdf (cheap, text only)
  read_image         - return one embedded image so the model can see it
  extract_images     - save every embedded image to a folder
  render_pdf_page    - render a PDF page as an image (charts, scans, vector art)
  insert_image_docx  - insert a picture into a Word document
  insert_image_pdf   - place a picture on a PDF page

Images sent to the model are downscaled (max_dim) to keep token use low.
Edits never overwrite the source unless output_path points at it.

Run: py -3.12 tools/docimage_mcp.py   (stdio transport)
"""

from __future__ import annotations

import io
from pathlib import Path

import pymupdf
from docx import Document
from docx.shared import Cm
from PIL import Image as PILImage

try:  # mcp 2.x renamed FastMCP to MCPServer
    from mcp.server.mcpserver import Image, MCPServer
except ImportError:  # mcp 1.x
    from mcp.server.fastmcp import FastMCP as MCPServer, Image

mcp = MCPServer("docimage")

MIN_SIDE = 24  # skip bullets, spacer pixels and other decorations
BLIP = "{http://schemas.openxmlformats.org/drawingml/2006/main}blip"
R_EMBED = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed"


def _path(p: str) -> Path:
    path = Path(p).expanduser().resolve()
    if not path.exists():
        raise FileNotFoundError(f"not found: {path}")
    return path


def _kind(path: Path) -> str:
    ext = path.suffix.lower()
    if ext == ".docx":
        return "docx"
    if ext == ".pdf":
        return "pdf"
    raise ValueError(f"unsupported file type {ext!r}; use .docx or .pdf")


def _docx_images(path: Path) -> list[dict]:
    """Embedded images in document order (body, then headers/footers)."""
    doc = Document(str(path))
    parts = [doc.part] + [
        s.part for sec in doc.sections for s in (sec.header, sec.footer) if s is not None
    ]
    seen, out = set(), []
    for part in parts:
        for blip in part.element.iter(BLIP):
            rid = blip.get(R_EMBED)
            if not rid or rid not in part.related_parts:
                continue
            img_part = part.related_parts[rid]
            key = img_part.partname
            if key in seen:
                continue
            seen.add(key)
            out.append({"blob": img_part.blob, "ext": Path(str(key)).suffix.lstrip(".") or "png", "page": None})
    return out


def _pdf_images(path: Path) -> list[dict]:
    out, seen = [], set()
    with pymupdf.open(str(path)) as doc:
        for pno, page in enumerate(doc, start=1):
            for info in page.get_images(full=True):
                xref = info[0]
                if xref in seen:
                    continue
                seen.add(xref)
                data = doc.extract_image(xref)
                if not data:
                    continue
                out.append({"blob": data["image"], "ext": data.get("ext", "png"), "page": pno})
    return out


def _images(path: Path) -> list[dict]:
    imgs = _docx_images(path) if _kind(path) == "docx" else _pdf_images(path)
    result = []
    for img in imgs:
        try:
            with PILImage.open(io.BytesIO(img["blob"])) as im:
                img["width"], img["height"] = im.size
        except Exception:
            img["width"] = img["height"] = None  # EMF/WMF etc.: keep, report unknown size
        if img["width"] and min(img["width"], img["height"]) < MIN_SIDE:
            continue
        result.append(img)
    return result


def _shrink(blob: bytes, max_dim: int) -> tuple[bytes, str]:
    with PILImage.open(io.BytesIO(blob)) as im:
        im = im.convert("RGBA") if im.mode in ("P", "LA") else im
        if max(im.size) > max_dim:
            im.thumbnail((max_dim, max_dim))
        buf = io.BytesIO()
        if im.mode in ("RGBA", "LA"):
            im.save(buf, "PNG", optimize=True)
            return buf.getvalue(), "png"
        im.convert("RGB").save(buf, "JPEG", quality=85)
        return buf.getvalue(), "jpeg"


def _default_output(src: Path) -> Path:
    return src.with_name(f"{src.stem}.edited{src.suffix}")


@mcp.tool()
def list_images(path: str) -> list[dict]:
    """List images embedded in a .docx or .pdf: index, page (PDF), format, pixel size.
    Use this first; it costs almost no tokens."""
    src = _path(path)
    return [
        {"index": i, "page": img["page"], "format": img["ext"], "width": img["width"], "height": img["height"]}
        for i, img in enumerate(_images(src))
    ]


@mcp.tool()
def read_image(path: str, index: int, max_dim: int = 1024) -> Image:
    """Return image number `index` (from list_images) so you can look at it.
    max_dim caps the longest side in pixels; lower it to save tokens."""
    imgs = _images(_path(path))
    if not 0 <= index < len(imgs):
        raise IndexError(f"index {index} out of range (0..{len(imgs) - 1})")
    data, fmt = _shrink(imgs[index]["blob"], max_dim)
    return Image(data=data, format=fmt)


@mcp.tool()
def extract_images(path: str, out_dir: str = "", original_quality: bool = True) -> list[str]:
    """Save every embedded image to out_dir (default: <file>_images next to the file).
    Returns the saved file paths."""
    src = _path(path)
    dest = Path(out_dir).expanduser().resolve() if out_dir else src.with_name(f"{src.stem}_images")
    dest.mkdir(parents=True, exist_ok=True)
    saved = []
    for i, img in enumerate(_images(src)):
        blob, ext = img["blob"], img["ext"]
        if not original_quality:
            blob, ext = _shrink(blob, 1600)
        page = f"_p{img['page']}" if img["page"] else ""
        target = dest / f"{src.stem}{page}_img{i:03d}.{ext}"
        target.write_bytes(blob)
        saved.append(str(target))
    return saved


@mcp.tool()
def render_pdf_page(path: str, page: int = 1, dpi: int = 100) -> Image:
    """Render a PDF page (1-based) as an image, including vector charts and scans.
    dpi 72-100 is usually readable; higher costs more tokens."""
    src = _path(path)
    if _kind(src) != "pdf":
        raise ValueError("render_pdf_page needs a .pdf")
    with pymupdf.open(str(src)) as doc:
        if not 1 <= page <= doc.page_count:
            raise IndexError(f"page {page} out of range (1..{doc.page_count})")
        pix = doc[page - 1].get_pixmap(dpi=dpi)
        return Image(data=pix.tobytes("png"), format="png")


@mcp.tool()
def insert_image_docx(
    docx_path: str,
    image_path: str,
    output_path: str = "",
    width_cm: float = 15.0,
    after_text: str = "",
    caption: str = "",
) -> str:
    """Insert a picture into a Word document.

    after_text: place it right after the first paragraph containing this text
                (default: end of document). caption: optional line under the picture.
    output_path: default <name>.edited.docx (the original is left untouched)."""
    src, img = _path(docx_path), _path(image_path)
    if _kind(src) != "docx":
        raise ValueError("insert_image_docx needs a .docx (convert .doc first)")
    doc = Document(str(src))

    anchor = None
    if after_text:
        anchor = next((p for p in doc.paragraphs if after_text in p.text), None)
        if anchor is None:
            raise ValueError(f"no paragraph contains {after_text!r}")

    pic_par = doc.add_paragraph()
    pic_par.add_run().add_picture(str(img), width=Cm(width_cm))
    new_pars = [pic_par]
    if caption:
        cap = doc.add_paragraph(caption)
        try:
            cap.style = doc.styles["Caption"]
        except KeyError:
            pass
        new_pars.append(cap)

    if anchor is not None:
        prev = anchor._p
        for par in new_pars:
            prev.addnext(par._p)
            prev = par._p

    out = Path(output_path).expanduser().resolve() if output_path else _default_output(src)
    doc.save(str(out))
    return str(out)


@mcp.tool()
def insert_image_pdf(
    pdf_path: str,
    image_path: str,
    page: int = 1,
    x: float = -1,
    y: float = -1,
    width: float = -1,
    height: float = -1,
    new_page: bool = False,
    output_path: str = "",
) -> str:
    """Place a picture on a PDF page (1-based). Coordinates are points (1/72 inch)
    from the top-left corner.

    Leave x/y/width/height at -1 to fit the image inside the page with 36pt margins.
    new_page=True appends a blank page (same size as the last page) and puts it there.
    output_path: default <name>.edited.pdf (the original is left untouched)."""
    src, img = _path(pdf_path), _path(image_path)
    if _kind(src) != "pdf":
        raise ValueError("insert_image_pdf needs a .pdf")
    with pymupdf.open(str(src)) as doc:
        if new_page:
            last = doc[-1].rect
            target = doc.new_page(width=last.width, height=last.height)
        else:
            if not 1 <= page <= doc.page_count:
                raise IndexError(f"page {page} out of range (1..{doc.page_count})")
            target = doc[page - 1]

        pr = target.rect
        if min(x, y, width, height) < 0:
            margin = 36
            rect = pymupdf.Rect(pr.x0 + margin, pr.y0 + margin, pr.x1 - margin, pr.y1 - margin)
            if x >= 0 and y >= 0 and width > 0:
                rect = pymupdf.Rect(x, y, x + width, pr.y1 - margin)
        else:
            rect = pymupdf.Rect(x, y, x + width, y + height)
        target.insert_image(rect, filename=str(img), keep_proportion=True)

        out = Path(output_path).expanduser().resolve() if output_path else _default_output(src)
        doc.save(str(out), garbage=3, deflate=True)
    return str(out)


if __name__ == "__main__":
    mcp.run()
